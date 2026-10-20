"""
DOCX Builder: Converts Markdown text into a styled, professional Word Document.
"""

import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def clean_xml_string(s):
    """Sanitizes text to be 100% XML compatible for python-docx."""
    if not isinstance(s, str):
        return str(s)
    # Remove XML-incompatible control characters
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f]', '', s)

def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_styled_document():
    doc = docx.Document()
    
    # Configure 1-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    return doc

def add_title_block(doc, title, subtitle, author, project_name):
    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(8)
    run_title = p_title.add_run(clean_xml_string(title))
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 58, 138) # Deep Navy

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(18)
    run_sub = p_sub.add_run(clean_xml_string(subtitle))
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(15)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(71, 85, 105) # Slate

    # Project metadata banner
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    meta_info = [
        ("Project System:", project_name, "Author & Engineer:", author),
        ("Target Platform:", "Python 3.11 / TensorFlow 2.15 / Streamlit", "Documentation Type:", "Master Technical Reference & Learning Guide")
    ]
    
    for row_idx, data in enumerate(meta_info):
        row = table.rows[row_idx]
        c0, c1 = row.cells[0], row.cells[1]
        
        c0.text = clean_xml_string(f"{data[0]} {data[1]}")
        c1.text = clean_xml_string(f"{data[2]} {data[3]}")
        
        for c in (c0, c1):
            set_cell_background(c, "F1F5F9")
            set_cell_margins(c, top=80, bottom=80, left=120, right=120)
            p = c.paragraphs[0]
            p.runs[0].font.name = 'Arial'
            p.runs[0].font.size = Pt(9.5)
            p.runs[0].font.color.rgb = RGBColor(51, 65, 85)

    doc.add_page_break()

def parse_markdown_to_docx(doc, md_text):
    lines = md_text.splitlines()
    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_table():
        nonlocal in_table, table_lines
        if not table_lines:
            return
        
        # Clean rows
        rows_data = []
        for line in table_lines:
            line_str = line.strip()
            if line_str.startswith("|") and line_str.endswith("|"):
                line_str = line_str[1:-1]
            cells = [clean_xml_string(c.strip()) for c in line_str.split("|")]
            # Skip separator line (e.g. :--- | :---)
            if all(re.match(r"^:?-+:?$", c) for c in cells if c):
                continue
            rows_data.append(cells)
            
        if not rows_data:
            table_lines = []
            in_table = False
            return
            
        col_count = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=col_count)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        for r_idx, r_cells in enumerate(rows_data):
            row = table.rows[r_idx]
            is_header = (r_idx == 0)
            bg_color = "1E3A8A" if is_header else ("F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            
            for c_idx in range(col_count):
                cell = row.cells[c_idx]
                text = r_cells[c_idx] if c_idx < len(r_cells) else ""
                cell.text = clean_xml_string(text)
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                if p.runs:
                    run = p.runs[0]
                    run.font.name = 'Arial'
                    run.font.size = Pt(9)
                    if is_header:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                    else:
                        run.font.color.rgb = RGBColor(30, 41, 59)
                        
        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_after = Pt(6)
        table_lines = []
        in_table = False

    def flush_code():
        nonlocal in_code_block, code_lines
        if not code_lines:
            in_code_block = False
            return
        
        code_text = clean_xml_string("\n".join(code_lines))
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        
        cell.text = code_text
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        if p.runs:
            run = p.runs[0]
            run.font.name = 'Consolas'
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(30, 41, 59)
            
        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_after = Pt(6)
        code_lines = []
        in_code_block = False

    for line in lines:
        stripped = line.strip()

        # Handle Code blocks
        if stripped.startswith("```"):
            if in_code_block:
                flush_code()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Handle Tables
        if "|" in line and (line.strip().startswith("|") or line.strip().endswith("|")):
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(line)
            continue
        else:
            if in_table:
                flush_table()

        # Empty lines
        if not stripped:
            continue

        # Headings
        if stripped.startswith("# "):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(22)
            h.paragraph_format.space_after = Pt(8)
            run = h.add_run(clean_xml_string(stripped[2:]))
            run.font.name = 'Arial'
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 64, 175) # Blue
            continue
        elif stripped.startswith("## "):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(6)
            run = h.add_run(clean_xml_string(stripped[3:]))
            run.font.name = 'Arial'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(37, 99, 235) # Light blue
            continue
        elif stripped.startswith("### "):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            run = h.add_run(clean_xml_string(stripped[4:]))
            run.font.name = 'Arial'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(51, 65, 85) # Slate
            continue
        elif stripped.startswith("#### "):
            h = doc.add_heading(level=4)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            run = h.add_run(clean_xml_string(stripped[5:]))
            run.font.name = 'Arial'
            run.font.size = Pt(10.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(71, 85, 105)
            continue

        # Blockquotes
        if stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(clean_xml_string(stripped[2:]))
            run.font.name = 'Arial'
            run.font.italic = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(71, 85, 105)
            continue

        # Bullet lists
        if stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            content = stripped[2:]
            _add_formatted_runs(p, content)
            continue

        # Numbered lists
        num_match = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if num_match:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            content = num_match.group(2)
            _add_formatted_runs(p, content)
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        _add_formatted_runs(p, stripped)

    if in_code_block:
        flush_code()
    if in_table:
        flush_table()

def _add_formatted_runs(paragraph, text):
    """Splits text by markdown bold and code ticks and adds formatted runs."""
    text = clean_xml_string(text)
    # Pattern to tokenize: **bold** and `code`
    tokens = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for token in tokens:
        if not token:
            continue
        token = clean_xml_string(token)
        if token.startswith("**") and token.endswith("**") and len(token) >= 4:
            run = paragraph.add_run(clean_xml_string(token[2:-2]))
            run.font.name = 'Arial'
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
        elif token.startswith("`") and token.endswith("`") and len(token) >= 2:
            run = paragraph.add_run(clean_xml_string(token[1:-1]))
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(185, 28, 28) # Reddish code text
        else:
            run = paragraph.add_run(clean_xml_string(token))
            run.font.name = 'Arial'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(30, 41, 59)
